from __future__ import annotations

from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
import threading

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote
import hashlib
import json
import math
import os
import re
import sys
import time
import xml.etree.ElementTree as ET

import pandas as pd
import requests
from dotenv import load_dotenv
from openpyxl import load_workbook

PIPELINE_DIR = Path(__file__).resolve().parents[1]
if str(PIPELINE_DIR) not in sys.path:
    sys.path.insert(0, str(PIPELINE_DIR))

from pipeline_elt import replace_raw_dataset_group


# ============================================================
# 1. PATH / ENV
# ============================================================

CURRENT_DIR = Path(__file__).resolve().parent

# 이 파일을 M3_data_pipeline/collector 안에 두는 현재 프로젝트 구조 기준
PROJECT_ROOT = (
    Path(__file__).resolve().parents[2]
    if len(Path(__file__).resolve().parents) >= 3
    else CURRENT_DIR
)

ENV_PATH = CURRENT_DIR / ".env"

if ENV_PATH.exists():
    load_dotenv(ENV_PATH, override=False)


def env_path(name: str, default: Path) -> Path:
    value = os.getenv(name)
    return Path(value).expanduser().resolve() if value else default.resolve()


CODELIST_PATH = env_path(
    "BICYCLE_CODELIST_PATH",
    CURRENT_DIR / "AccidentHazard_CodeList.xlsx",
)

RAW_DIR = env_path(
    "BICYCLE_RAW_DIR",
    PROJECT_ROOT / "data" / "raw" / "koroad_bicycle",
)

CHECKPOINT_PATH = RAW_DIR / "koroad_bicycle_checkpoint.jsonl"


# ============================================================
# 2. API / COLLECTION SETTINGS
# ============================================================

API_URL = (
    "https://apis.data.go.kr/B552061/"
    "frequentzoneBicycle/getRestFrequentzoneBicycle"
)

RAW_API_KEY = os.getenv("BICYCLE_ACCIDENT_API_KEY", "").strip()

# 포털에서 URL Encode 상태로 준 키를 한 번 복원하고,
# requests가 query string을 만들 때 한 번만 인코딩하게 한다.
API_KEY = unquote(RAW_API_KEY)

MIN_YEAR = 2012
MAX_YEAR = 2024
NUM_OF_ROWS = 100
MAX_RETRIES = 3
RETRY_WAIT_SECONDS = (3, 5)
REQUEST_TIMEOUT = (10, 45)
REQUEST_DELAY_SECONDS = 0.05

# 병렬 수집 설정
MAX_WORKERS = 6

_thread_local = threading.local()

API_COLUMNS = [
    "afos_fid",
    "afos_id",
    "bjd_cd",
    "spot_cd",
    "sido_sgg_nm",
    "spot_nm",
    "occrrnc_cnt",
    "caslt_cnt",
    "dth_dnv_cnt",
    "se_dnv_cnt",
    "sl_dnv_cnt",
    "wnd_dnv_cnt",
    "lo_crd",
    "la_crd",
    "geom_json",
]

REQUEST_METADATA_COLUMNS = [
    "request_year",
    "request_sido",
    "request_gugun",
    "request_province_name",
    "request_district_name",
    "expected_afos_id",
    "request_page_no",
    "collected_at",
]

# ============================================================
# 3. EXCEPTIONS
# ============================================================


class PipelineError(RuntimeError):
    pass


class TransientApiError(PipelineError):
    """재시도할 수 있는 일시적 API 오류."""


class PermanentApiError(PipelineError):
    """재시도해도 해결되지 않는 요청/응답 오류."""


class UnsupportedScopeError(PermanentApiError):
    """해당 연도에 존재하지 않거나 지원되지 않는 행정구역 조합."""


class GlobalServiceError(PermanentApiError):
    """인증키, 사용승인 또는 호출 한도와 관련된 전역 서비스 오류."""


# ============================================================
# 4. CODE LIST
# ============================================================


def normalize_name(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def name_key(value) -> str:
    return re.sub(r"\s+", "", normalize_name(value))


def load_request_scopes(code_list_path: Path):
    if not code_list_path.exists():
        raise FileNotFoundError(
            "요청변수 코드표를 찾을 수 없습니다.\n"
            f"확인 경로: {code_list_path}"
        )

    workbook = load_workbook(
        code_list_path,
        read_only=False,
        data_only=True,
    )

    required_sheets = {
        "serachYearCd 요청값",
        "Sido 요청값",
        "Gugun 요청값",
    }
    missing_sheets = required_sheets.difference(workbook.sheetnames)

    if missing_sheets:
        raise PipelineError(
            "코드표 필수 시트가 없습니다: "
            + ", ".join(sorted(missing_sheets))
        )

    year_sheet = workbook["serachYearCd 요청값"]
    sido_sheet = workbook["Sido 요청값"]
    gugun_sheet = workbook["Gugun 요청값"]

    sido_by_name: dict[str, str] = {}

    for row in range(2, sido_sheet.max_row + 1):
        sido_name = sido_sheet.cell(row, 1).value
        sido_code = sido_sheet.cell(row, 2).value

        if sido_name is None or sido_code is None:
            continue

        sido_by_name[name_key(sido_name)] = f"{int(sido_code):02d}"

    afos_id_by_year: dict[int, str] = {}
    current_category = ""

    for row in range(2, year_sheet.max_row + 1):
        category = year_sheet.cell(row, 1).value
        label = year_sheet.cell(row, 2).value
        dataset_id = year_sheet.cell(row, 3).value

        if category is not None:
            current_category = normalize_name(category)

        if current_category != "자전거 교통사고 다발지역":
            continue

        match = re.match(r"(\d{2})년", normalize_name(label))

        if not match or dataset_id is None:
            continue

        year = 2000 + int(match.group(1))

        if MIN_YEAR <= year <= MAX_YEAR:
            afos_id_by_year[year] = str(int(dataset_id))

    expected_years = list(range(MIN_YEAR, MAX_YEAR + 1))

    if sorted(afos_id_by_year) != expected_years:
        raise PipelineError(
            "자전거 제공연도 코드가 2012~2024와 일치하지 않습니다."
        )

    gugun_by_province: dict[str, list[dict[str, str]]] = defaultdict(list)
    current_province = None

    for row in range(2, gugun_sheet.max_row + 1):
        province = gugun_sheet.cell(row, 1).value
        district = gugun_sheet.cell(row, 2).value
        gugun_code = gugun_sheet.cell(row, 3).value

        if province is not None:
            current_province = normalize_name(province)

        if district is None or gugun_code is None:
            continue

        if current_province is None:
            raise PipelineError(
                f"Gugun 요청값 {row}행에 시도 구분이 없습니다."
            )

        gugun_by_province[current_province].append(
            {
                "district_name": normalize_name(district),
                "guGun": f"{int(gugun_code):03d}",
            }
        )

    def resolve_sido(year: int, province: str) -> str:
        province_key = name_key(province)

        if province_key == name_key("강원특별자치도"):
            source_name = (
                "강원도(구)" if year <= 2022 else "강원특별자치도"
            )
            return sido_by_name[name_key(source_name)]

        if province_key == name_key("전북특별자치도"):
            source_name = (
                "전라북도(구)" if year <= 2022 else "전북특별자치도"
            )
            return sido_by_name[name_key(source_name)]

        if province_key not in sido_by_name:
            raise PipelineError(
                f"시도 코드 매핑을 찾을 수 없습니다: {province}"
            )

        return sido_by_name[province_key]

    scopes = []

    for year in expected_years:
        for province, districts in gugun_by_province.items():
            sido_code = resolve_sido(year, province)

            # 공식 제공기간 2012~2024에서는 구 광주/전남 코드(29/46)를 사용한다.
            if sido_code == "12":
                continue

            for district in districts:
                scopes.append(
                    {
                        "searchYearCd": str(year),
                        "siDo": sido_code,
                        "guGun": district["guGun"],
                        "province_name": province,
                        "district_name": district["district_name"],
                        "expected_afos_id": afos_id_by_year[year],
                    }
                )

    if len(scopes) != 3510:
        raise PipelineError(
            "예상 요청범위 3,510개와 다릅니다: "
            f"{len(scopes):,}개"
        )

    return scopes, afos_id_by_year


# ============================================================
# 5. API RESPONSE PARSING
# ============================================================


def safe_response_preview(response: requests.Response, limit: int = 500) -> str:
    preview = response.text[:limit]

    for secret in (RAW_API_KEY, API_KEY):
        if secret:
            preview = preview.replace(secret, "[API_KEY]")

    return repr(preview)


def normalize_items(items_container) -> list[dict]:
    if not items_container:
        return []

    if isinstance(items_container, dict):
        items = items_container.get("item", [])
    else:
        items = items_container

    if not items:
        return []

    if isinstance(items, dict):
        items = [items]

    if not isinstance(items, list):
        raise PermanentApiError(
            "API items.item 구조가 예상과 다릅니다."
        )

    if not all(isinstance(item, dict) for item in items):
        raise PermanentApiError(
            "API item 중 객체가 아닌 값이 있습니다."
        )

    return items


def parse_json_payload(data: dict) -> dict:
    # KoROAD 현재 형식은 최상위에 resultCode/items/totalCount가 있다.
    # response/header/body 형식도 방어적으로 지원한다.
    if isinstance(data.get("response"), dict):
        response_data = data["response"]
        header = response_data.get("header", {})
        body = response_data.get("body", {})

        result_code = header.get("resultCode")
        result_msg = header.get("resultMsg")
        total_count = body.get("totalCount", 0)
        num_of_rows = body.get("numOfRows")
        page_no = body.get("pageNo")
        items = normalize_items(body.get("items"))
    else:
        result_code = data.get("resultCode")
        result_msg = data.get("resultMsg")
        total_count = data.get("totalCount", 0)
        num_of_rows = data.get("numOfRows")
        page_no = data.get("pageNo")
        items = normalize_items(data.get("items"))

    return {
        "result_code": (
            "" if result_code is None else str(result_code).strip()
        ),
        "result_msg": (
            "" if result_msg is None else str(result_msg).strip()
        ),
        "total_count": total_count,
        "num_of_rows": num_of_rows,
        "page_no": page_no,
        "items": items,
    }


def parse_xml_payload(xml_text: str) -> dict:
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError as exc:
        raise TransientApiError(
            "API 응답을 JSON 또는 XML로 해석할 수 없습니다."
        ) from exc

    items = []

    for item_node in root.findall(".//items/item"):
        item = {}

        for child in list(item_node):
            item[child.tag] = child.text

        items.append(item)

    return {
        "result_code": str(root.findtext(".//resultCode") or "").strip(),
        "result_msg": str(root.findtext(".//resultMsg") or "").strip(),
        "total_count": root.findtext(".//totalCount") or 0,
        "num_of_rows": root.findtext(".//numOfRows"),
        "page_no": root.findtext(".//pageNo"),
        "items": items,
    }


def parse_api_response(response: requests.Response) -> dict:
    try:
        data = response.json()
    except ValueError:
        return parse_xml_payload(response.text)

    if not isinstance(data, dict):
        raise TransientApiError(
            "API JSON 최상위 구조가 객체가 아닙니다."
        )

    return parse_json_payload(data)


def as_nonnegative_int(value, field_name: str) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise PermanentApiError(
            f"{field_name}를 정수로 변환할 수 없습니다: {value!r}"
        ) from exc

    if result < 0:
        raise PermanentApiError(
            f"{field_name}가 음수입니다: {result}"
        )

    return result


# ============================================================
# 6. API REQUEST / RETRY / PAGINATION
# ============================================================


def request_page_once(
    session: requests.Session,
    scope: dict,
    page_no: int,
) -> dict:
    params = {
        "serviceKey": API_KEY,
        "searchYearCd": scope["searchYearCd"],
        "siDo": scope["siDo"],
        "guGun": scope["guGun"],
        "type": "json",
        "numOfRows": str(NUM_OF_ROWS),
        "pageNo": str(page_no),
    }

    try:
        response = session.get(
            API_URL,
            params=params,
            timeout=REQUEST_TIMEOUT,
        )
    except (requests.Timeout, requests.ConnectionError) as exc:
        raise TransientApiError(str(exc)) from exc
    except requests.RequestException as exc:
        raise TransientApiError(str(exc)) from exc

    if response.status_code == 429:
        response_preview = safe_response_preview(response)

        # data.go.kr 일일 서비스 요청 한도 초과
        if (
            "LIMITED_NUMBER_OF_SERVICE_REQUESTS_EXCEEDS_ERROR" in response.text
            or '"returnReasonCode": "22"' in response.text
            or '"returnReasonCode":"22"' in response.text
        ):
            raise GlobalServiceError(
                "data.go.kr 일일 서비스 요청 한도를 초과했습니다. "
                f"HTTP 429: {response_preview}"
            )

        raise TransientApiError(
            f"HTTP 429: {response_preview}"
        )

    if response.status_code >= 500:
        raise TransientApiError(
            f"HTTP {response.status_code}: "
            f"{safe_response_preview(response)}"
        )

    if response.status_code != 200:
        raise PermanentApiError(
            f"HTTP {response.status_code}: "
            f"{safe_response_preview(response)}"
        )

    parsed = parse_api_response(response)
    result_code = parsed["result_code"]
    result_msg = parsed["result_msg"]

    if result_code == "03":
        parsed["total_count"] = 0
        parsed["items"] = []
        return parsed

    if result_code == "99":
        raise TransientApiError(
            f"API 오류 99: {result_msg}"
        )

    # 코드표는 폐지·신설 행정구역을 모두 포함한다. 인증키와 기본 요청이
    # 정상임을 사전 점검한 뒤에는, 개별 연도에 유효하지 않은 조합에서
    # 반환되는 10을 수집 전체의 장애로 보지 않고 별도로 건너뛴다.
    if result_code == "10":
        raise UnsupportedScopeError(
            f"지원되지 않는 연도/행정구역 조합: {result_msg}"
        )

    if result_code in {"30", "31", "32"}:
        raise GlobalServiceError(
            f"API 전역 서비스 오류 {result_code}: {result_msg}"
        )

    if result_code not in {"00", "0", "0000"}:
        raise PermanentApiError(
            f"API 오류 {result_code or '[없음]'}: {result_msg}"
        )

    parsed["total_count"] = as_nonnegative_int(
        parsed["total_count"],
        "totalCount",
    )
    return parsed


def request_page_with_retry(
    session: requests.Session,
    scope: dict,
    page_no: int,
) -> dict:
    last_error = None

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            return request_page_once(session, scope, page_no)
        except PermanentApiError:
            raise
        except TransientApiError as exc:
            last_error = exc

            if attempt < MAX_RETRIES:
                wait_seconds = RETRY_WAIT_SECONDS[attempt - 1]
                print(
                    f"    페이지 {page_no} 요청 실패 "
                    f"({attempt}/{MAX_RETRIES}) → "
                    f"{wait_seconds}초 후 재시도"
                )
                time.sleep(wait_seconds)

    raise TransientApiError(
        f"페이지 {page_no} 요청이 {MAX_RETRIES}회 모두 실패했습니다: "
        f"{last_error}"
    ) from last_error


def attach_request_metadata(
    item: dict,
    scope: dict,
    page_no: int,
    collected_at: str,
) -> dict:
    record = dict(item)
    record.update(
        {
            "request_year": scope["searchYearCd"],
            "request_sido": scope["siDo"],
            "request_gugun": scope["guGun"],
            "request_province_name": scope["province_name"],
            "request_district_name": scope["district_name"],
            "expected_afos_id": scope["expected_afos_id"],
            "request_page_no": page_no,
            "collected_at": collected_at,
        }
    )
    return record


def scope_key(scope: dict) -> str:
    return "|".join(
        [
            str(scope["searchYearCd"]),
            str(scope["siDo"]),
            str(scope["guGun"]),
        ]
    )


def checkpoint_meta(scopes: list[dict]) -> dict:
    joined_keys = "\n".join(scope_key(scope) for scope in scopes)
    fingerprint = hashlib.sha256(
        joined_keys.encode("utf-8")
    ).hexdigest()

    return {
        "type": "meta",
        "version": 1,
        "min_year": MIN_YEAR,
        "max_year": MAX_YEAR,
        "scope_count": len(scopes),
        "scope_fingerprint": fingerprint,
    }


def rewrite_checkpoint(meta: dict, entries: dict[str, dict]) -> None:
    temp_path = CHECKPOINT_PATH.with_suffix(".tmp.jsonl")

    with temp_path.open("w", encoding="utf-8", newline="\n") as file:
        file.write(json.dumps(meta, ensure_ascii=False) + "\n")

        for entry in entries.values():
            file.write(
                json.dumps(
                    entry,
                    ensure_ascii=False,
                    separators=(",", ":"),
                    default=str,
                )
                + "\n"
            )

        file.flush()
        os.fsync(file.fileno())

    os.replace(temp_path, CHECKPOINT_PATH)


def load_checkpoint(scopes: list[dict]) -> tuple[dict, dict[str, dict]]:
    expected_meta = checkpoint_meta(scopes)

    if not CHECKPOINT_PATH.exists():
        rewrite_checkpoint(expected_meta, {})
        return expected_meta, {}

    valid_scope_keys = {scope_key(scope) for scope in scopes}
    entries: dict[str, dict] = {}
    actual_meta = None
    invalid_line_found = False

    with CHECKPOINT_PATH.open("r", encoding="utf-8") as file:
        for line_number, line in enumerate(file, start=1):
            if not line.strip():
                continue

            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                invalid_line_found = True
                print(
                    f"체크포인트 {line_number}행이 불완전하여 "
                    "마지막 정상 범위부터 복구합니다."
                )
                continue

            if record.get("type") == "meta":
                actual_meta = record
                continue

            if record.get("type") != "scope":
                continue

            key = record.get("key")

            if key in valid_scope_keys:
                entries[key] = record

    if actual_meta is None:
        raise PipelineError(
            f"체크포인트 메타데이터가 없습니다: {CHECKPOINT_PATH}"
        )

    fields_to_compare = [
        "version",
        "min_year",
        "max_year",
        "scope_count",
        "scope_fingerprint",
    ]

    if any(
        actual_meta.get(field) != expected_meta.get(field)
        for field in fields_to_compare
    ):
        raise PipelineError(
            "기존 체크포인트의 코드표/연도 범위가 현재 실행과 다릅니다.\n"
            f"확인 파일: {CHECKPOINT_PATH}"
        )

    if invalid_line_found:
        rewrite_checkpoint(expected_meta, entries)

    return expected_meta, entries


def append_checkpoint_entry(entry: dict) -> None:
    with CHECKPOINT_PATH.open("a", encoding="utf-8", newline="\n") as file:
        file.write(
            json.dumps(
                entry,
                ensure_ascii=False,
                separators=(",", ":"),
                default=str,
            )
            + "\n"
        )
        file.flush()
        os.fsync(file.fileno())


def make_checkpoint_entry(
    scope: dict,
    status: str,
    records: list[dict] | None = None,
    reason: str | None = None,
) -> dict:
    return {
        "type": "scope",
        "key": scope_key(scope),
        "status": status,
        "scope": scope,
        "records": records or [],
        "reason": reason,
        "saved_at": datetime.now()
        .astimezone()
        .isoformat(timespec="seconds"),
    }


def remove_checkpoint_after_success() -> None:
    if CHECKPOINT_PATH.exists():
        CHECKPOINT_PATH.unlink()


def get_thread_session():
    """
    병렬 수집 worker별 requests.Session을 생성하고 재사용한다.
    Session 객체를 여러 thread가 공유하지 않는다.
    """

    if not hasattr(_thread_local, "session"):
        session = requests.Session()
        session.headers.update(
            {
                "Accept": "application/json",
                "User-Agent": "WooSimWoonKka-KoROAD-Collector/1.0",
            }
        )
        _thread_local.session = session

    return _thread_local.session


def collect_scope_worker(scope, collected_at):
    """
    하나의 요청범위를 worker thread에서 수집한다.
    scope 내부 pagination은 기존 collect_scope()가 순차 처리한다.
    """

    session = get_thread_session()

    try:
        return collect_scope(
            session,
            scope,
            collected_at,
        )
    finally:
        # 같은 worker가 다음 scope의 첫 요청을 즉시 연속 호출하지 않도록
        # worker별 pacing을 적용한다. 전역 lock은 사용하지 않는다.
        time.sleep(REQUEST_DELAY_SECONDS)


def collect_scope(
    session: requests.Session,
    scope: dict,
    collected_at: str,
) -> list[dict]:
    first_page = request_page_with_retry(session, scope, 1)
    total_count = first_page["total_count"]

    if total_count == 0:
        return []

    total_pages = math.ceil(total_count / NUM_OF_ROWS)
    records = [
        attach_request_metadata(item, scope, 1, collected_at)
        for item in first_page["items"]
    ]

    for page_no in range(2, total_pages + 1):
        # 직전 페이지와 다음 페이지 요청 사이의 pacing.
        time.sleep(REQUEST_DELAY_SECONDS)
        page = request_page_with_retry(session, scope, page_no)

        records.extend(
            attach_request_metadata(item, scope, page_no, collected_at)
            for item in page["items"]
        )

    if len(records) != total_count:
        raise PermanentApiError(
            "페이지 수집 건수 불일치: "
            f"totalCount={total_count:,} / 수집={len(records):,}"
        )

    return records


def save_failure_log(failures: list[dict], run_id: str) -> Path:
    failure_path = RAW_DIR / f"koroad_bicycle_failures_{run_id}.csv"
    temp_path = failure_path.with_suffix(".tmp.csv")

    pd.DataFrame(failures).to_csv(
        temp_path,
        index=False,
        encoding="utf-8-sig",
    )
    os.replace(temp_path, failure_path)
    return failure_path


def save_skipped_scope_log(skipped_scopes: list[dict], run_id: str) -> Path:
    skipped_path = RAW_DIR / f"koroad_bicycle_skipped_scopes_{run_id}.csv"
    temp_path = skipped_path.with_suffix(".tmp.csv")

    pd.DataFrame(skipped_scopes).to_csv(
        temp_path,
        index=False,
        encoding="utf-8-sig",
    )
    os.replace(temp_path, skipped_path)
    return skipped_path


def records_in_scope_order(
    scopes: list[dict],
    checkpoint_entries: dict[str, dict],
) -> list[dict]:
    """완료 순서와 무관하게 원래 scope 순서로 success record를 조립한다."""

    records: list[dict] = []

    for scope in scopes:
        entry = checkpoint_entries.get(scope_key(scope))

        if entry and entry.get("status") == "success":
            records.extend(entry.get("records") or [])

    return records


def print_collection_metrics(
    *,
    total_scopes: int,
    resumed_scopes: int,
    api_target_scopes: int,
    checkpoint_entries: dict[str, dict],
    failed_scopes: int,
    started_at: float,
) -> None:
    status_counts = {
        status: sum(
            entry.get("status") == status
            for entry in checkpoint_entries.values()
        )
        for status in ("success", "no_data", "unsupported")
    }
    wall_clock_minutes = (time.perf_counter() - started_at) / 60

    print()
    print("[수집 실행 요약]")
    print(f"전체 scope      : {total_scopes:,}개")
    print(f"재개 scope      : {resumed_scopes:,}개")
    print(f"API 대상 scope  : {api_target_scopes:,}개")
    print(f"success         : {status_counts['success']:,}개")
    print(f"no_data         : {status_counts['no_data']:,}개")
    print(f"unsupported     : {status_counts['unsupported']:,}개")
    print(f"failed          : {failed_scopes:,}개")
    print(f"MAX_WORKERS     : {MAX_WORKERS}")
    wall_clock_seconds = round(wall_clock_minutes * 60)
    print(f"wall-clock      : {wall_clock_seconds // 60}분 {wall_clock_seconds % 60}초")


def collect_all_scopes(scopes: list[dict], run_id: str) -> pd.DataFrame:
    collection_started = time.perf_counter()
    collected_at = datetime.now().astimezone().isoformat(timespec="seconds")
    failures: list[dict] = []
    skipped_scopes: list[dict] = []
    no_data_count = 0

    _, checkpoint_entries = load_checkpoint(scopes)
    resumed_scope_count = len(checkpoint_entries)

    for entry in checkpoint_entries.values():
        status = entry.get("status")

        if status == "no_data":
            no_data_count += 1
        elif status == "unsupported":
            saved_scope = entry.get("scope") or {}
            skipped_scopes.append(
                {
                    "searchYearCd": saved_scope.get("searchYearCd"),
                    "siDo": saved_scope.get("siDo"),
                    "guGun": saved_scope.get("guGun"),
                    "province_name": saved_scope.get("province_name"),
                    "district_name": saved_scope.get("district_name"),
                    "reason": entry.get("reason"),
                }
            )

    if checkpoint_entries:
        print(
            "체크포인트 재개: "
            f"완료된 요청 {len(checkpoint_entries):,}개를 건너뜁니다."
        )
        print(f"체크포인트     : {CHECKPOINT_PATH}")
        print()

    session = requests.Session()
    session.headers.update(
        {
            "Accept": "application/json",
            "User-Agent": "WooSimWoonKka-KoROAD-Collector/1.0",
        }
    )

    total_scopes = len(scopes)

    try:
        # 대량 수집 전에 이미 성공이 확인된 공식 예제 조합으로 인증키,
        # URL, JSON 형식과 기본 파라미터를 검증한다. 이 단계가 실패하면
        # resultCode=10을 행정구역 예외로 숨기지 않고 즉시 중단한다.
        preflight_scope = {
            "searchYearCd": "2023",
            "siDo": "11",
            "guGun": "680",
            "province_name": "서울특별시",
            "district_name": "강남구",
            "expected_afos_id": "2024046",
        }
        try:
            preflight = request_page_with_retry(
                session,
                preflight_scope,
                1,
            )
        except GlobalServiceError as exc:
            failure_path = save_failure_log(
                [
                    {
                        "searchYearCd": "2023",
                        "siDo": "11",
                        "guGun": "680",
                        "province_name": "서울특별시",
                        "district_name": "강남구",
                        "error_type": type(exc).__name__,
                        "error_message": str(exc),
                        "failed_at": datetime.now()
                        .astimezone()
                        .isoformat(timespec="seconds"),
                    }
                ],
                run_id,
            )
            raise PipelineError(
                "사전 API 점검에서 인증키 사용승인 또는 호출 한도 "
                "오류가 발생했습니다. 대량 수집을 시작하지 않습니다.\n"
                f"다음 실행은 체크포인트부터 이어집니다: "
                f"{CHECKPOINT_PATH}\n"
                f"오류 기록: {failure_path}"
            ) from exc
        print(
            "사전 API 점검 : 정상 "
            f"(2023 / 서울 / 강남구, "
            f"{preflight['total_count']:,}건)"
        )
        print()

        # =====================================================
        # 전국·전연도 요청범위 병렬 수집
        #
        # - API 호출만 worker thread에서 수행
        # - scope 내부 pagination은 기존 collect_scope()에서 순차 처리
        # - checkpoint 기록 / 결과 병합은 main thread에서만 수행
        # - worker별 requests.Session 사용
        # =====================================================

        pending_scopes = [
            (index, scope)
            for index, scope in enumerate(scopes, start=1)
            if scope_key(scope) not in checkpoint_entries
        ]

        print(f"병렬 Worker     : {MAX_WORKERS}")
        print(f"남은 요청 범위 : {len(pending_scopes):,}개")
        print()

        global_service_error = None
        stop_submitting = False
        pending_scope_iterator = iter(pending_scopes)
        future_to_scope = {}
        executor = ThreadPoolExecutor(max_workers=MAX_WORKERS)

        def submit_next_scope() -> bool:
            try:
                index, scope = next(pending_scope_iterator)
            except StopIteration:
                return False

            future = executor.submit(
                collect_scope_worker,
                scope,
                collected_at,
            )
            future_to_scope[future] = (index, scope)
            return True

        try:
            for _ in range(min(MAX_WORKERS, len(pending_scopes))):
                submit_next_scope()

            while future_to_scope:
                done_futures, _ = wait(
                    tuple(future_to_scope),
                    return_when=FIRST_COMPLETED,
                )
                refill_slots = 0

                for future in done_futures:
                    index, scope = future_to_scope.pop(future)

                    if future.cancelled():
                        continue

                    refill_slots += 1

                    key = scope_key(scope)

                    label = (
                        f"{scope['searchYearCd']} / "
                        f"{scope['province_name']}({scope['siDo']}) / "
                        f"{scope['district_name']}({scope['guGun']})"
                    )

                    try:
                        records = future.result()

                        # -------------------------------------
                        # 정상 데이터
                        # -------------------------------------

                        if records:

                            entry = make_checkpoint_entry(
                                scope,
                                "success",
                                records=records,
                            )

                            # checkpoint는 main thread에서만 기록
                            append_checkpoint_entry(entry)

                            checkpoint_entries[key] = entry

                            result_text = f"{len(records):,}건"

                        # -------------------------------------
                        # 정상 응답이지만 데이터 없음
                        # -------------------------------------

                        else:

                            entry = make_checkpoint_entry(
                                scope,
                                "no_data",
                            )

                            append_checkpoint_entry(entry)

                            checkpoint_entries[key] = entry
                            no_data_count += 1

                            result_text = "데이터 없음"

                        print(
                            f"[{index:04d}/{total_scopes:04d}] "
                            f"{label} → {result_text}"
                        )

                    # -----------------------------------------
                    # 해당 연도에 존재하지 않는 행정구역 조합
                    # -----------------------------------------

                    except UnsupportedScopeError as exc:

                        entry = make_checkpoint_entry(
                            scope,
                            "unsupported",
                            reason=str(exc),
                        )

                        append_checkpoint_entry(entry)

                        checkpoint_entries[key] = entry

                        skipped_scopes.append(
                            {
                                "searchYearCd": scope["searchYearCd"],
                                "siDo": scope["siDo"],
                                "guGun": scope["guGun"],
                                "province_name": scope[
                                    "province_name"
                                ],
                                "district_name": scope[
                                    "district_name"
                                ],
                                "reason": str(exc),
                            }
                        )

                        print(
                            f"[{index:04d}/{total_scopes:04d}] "
                            f"{label} → "
                            "해당 연도 미지원 조합, 건너뜀"
                        )

                    # -----------------------------------------
                    # 인증키 / 사용승인 / 호출한도 등 전역 오류
                    # -----------------------------------------

                    except GlobalServiceError as exc:

                        if global_service_error is None:
                            global_service_error = {
                                "searchYearCd": scope["searchYearCd"],
                                "siDo": scope["siDo"],
                                "guGun": scope["guGun"],
                                "province_name": scope[
                                    "province_name"
                                ],
                                "district_name": scope[
                                    "district_name"
                                ],
                                "error_type": type(exc).__name__,
                                "error_message": str(exc),
                                "failed_at": (
                                    datetime.now()
                                    .astimezone()
                                    .isoformat(timespec="seconds")
                                ),
                            }

                        print(
                            f"[{index:04d}/{total_scopes:04d}] "
                            f"{label} → 전역 서비스 오류: {exc}"
                        )

                        stop_submitting = True

                        # 아직 실행되지 않은 future는 취소한다.
                        # 실행 중인 최대 MAX_WORKERS개만 종료될 수 있다.
                        for pending_future in future_to_scope:
                            pending_future.cancel()

                    # -----------------------------------------
                    # 개별 scope 실패
                    # -----------------------------------------

                    except Exception as exc:

                        failures.append(
                            {
                                "searchYearCd": scope["searchYearCd"],
                                "siDo": scope["siDo"],
                                "guGun": scope["guGun"],
                                "province_name": scope[
                                    "province_name"
                                ],
                                "district_name": scope[
                                    "district_name"
                                ],
                                "error_type": type(exc).__name__,
                                "error_message": str(exc),
                                "failed_at": (
                                    datetime.now()
                                    .astimezone()
                                    .isoformat(timespec="seconds")
                                ),
                            }
                        )

                        print(
                            f"[{index:04d}/{total_scopes:04d}] "
                            f"{label} → 실패: {exc}"
                        )

                if not stop_submitting:
                    for _ in range(refill_slots):
                        if not submit_next_scope():
                            break

        except KeyboardInterrupt:
            stop_submitting = True

            for future in future_to_scope:
                future.cancel()

            executor.shutdown(
                wait=False,
                cancel_futures=True,
            )
            print_collection_metrics(
                total_scopes=total_scopes,
                resumed_scopes=resumed_scope_count,
                api_target_scopes=len(pending_scopes),
                checkpoint_entries=checkpoint_entries,
                failed_scopes=len(failures),
                started_at=collection_started,
            )
            print(
                "수집 중단: 신규 scope 제출을 중단하고 "
                "미실행 future를 취소했습니다."
            )
            raise

        except BaseException:
            for future in future_to_scope:
                future.cancel()

            executor.shutdown(
                wait=True,
                cancel_futures=True,
            )
            raise

        else:
            executor.shutdown(
                wait=True,
                cancel_futures=True,
            )

        print_collection_metrics(
            total_scopes=total_scopes,
            resumed_scopes=resumed_scope_count,
            api_target_scopes=len(pending_scopes),
            checkpoint_entries=checkpoint_entries,
            failed_scopes=(
                len(failures)
                + (1 if global_service_error is not None else 0)
            ),
            started_at=collection_started,
        )

        # =====================================================
        # 전역 서비스 오류 처리
        # =====================================================

        if global_service_error is not None:

            failure_path = save_failure_log(
                [global_service_error],
                run_id,
            )

            raise PipelineError(
                "인증키 사용승인 또는 일일 호출 한도와 관련된 "
                "전역 오류가 발생해 수집을 중단합니다.\n"
                f"완료 요청: {len(checkpoint_entries):,}/"
                f"{total_scopes:,}\n"
                f"다음 실행은 체크포인트부터 이어집니다: "
                f"{CHECKPOINT_PATH}\n"
                f"오류 기록: {failure_path}"
            )

    finally:
        session.close()

    print()
    print(
        "데이터 있는 범위 : "
        f"{len(checkpoint_entries) - no_data_count - len(skipped_scopes):,}개"
    )
    print(f"데이터 없는 범위 : {no_data_count:,}개")
    print(f"연도 미지원 조합 : {len(skipped_scopes):,}개")
    print(f"실패 범위        : {len(failures):,}개")
    print(f"완료 체크포인트  : {len(checkpoint_entries):,}/{total_scopes:,}개")

    if skipped_scopes:
        skipped_path = save_skipped_scope_log(skipped_scopes, run_id)
        print(f"건너뛴 조합 목록 : {skipped_path}")

    if failures:
        failure_path = save_failure_log(failures, run_id)
        raise PipelineError(
            "일부 요청이 실패하여 RAW DB를 교체하지 않습니다.\n"
            f"실패 목록: {failure_path}"
        )

    if len(checkpoint_entries) != total_scopes:
        raise PipelineError(
            "완료된 요청범위가 전체 요청범위와 다르므로 RAW DB를 "
            "교체하지 않습니다.\n"
            f"완료={len(checkpoint_entries):,} / 전체={total_scopes:,}\n"
            f"체크포인트: {CHECKPOINT_PATH}"
        )

    all_records = records_in_scope_order(
        scopes,
        checkpoint_entries,
    )

    if not all_records:
        raise PipelineError(
            "전체 수집 결과가 0건이므로 기존 RAW DB를 교체하지 않습니다."
        )

    raw_df = pd.DataFrame(all_records)

    collected_years = sorted(
        pd.to_numeric(raw_df["request_year"], errors="coerce")
        .dropna()
        .astype(int)
        .unique()
        .tolist()
    )
    expected_years = list(range(MIN_YEAR, MAX_YEAR + 1))

    if collected_years != expected_years:
        raise PipelineError(
            "수집 데이터의 연도 범위가 불완전하여 RAW DB를 "
            "교체하지 않습니다.\n"
            f"예상={expected_years} / 실제={collected_years}"
        )

    missing_columns = set(API_COLUMNS).difference(raw_df.columns)

    if missing_columns:
        raise PipelineError(
            "API 필수 컬럼이 없습니다: "
            + ", ".join(sorted(missing_columns))
        )

    ordered_columns = API_COLUMNS + REQUEST_METADATA_COLUMNS
    extra_columns = [
        column
        for column in raw_df.columns
        if column not in ordered_columns
    ]

    return raw_df.reindex(columns=ordered_columns + extra_columns)


# ============================================================
# 7. RAW DB ATOMIC REPLACEMENT
# ============================================================


def replace_raw_snapshot(raw_df: pd.DataFrame) -> None:
    replace_raw_dataset_group(
        {"koroad_bicycle_accident_hotspots": raw_df}
    )


# ============================================================
# 8. OUTPUT
# ============================================================


def display_first_record(raw_df: pd.DataFrame) -> None:
    first = raw_df.iloc[0]

    def display(column):
        value = first.get(column)
        return "NULL" if pd.isna(value) else value

    print()
    print("[첫 번째 수집 데이터]")
    print(f"지역         : {display('sido_sgg_nm')}")
    print(f"지점         : {display('spot_nm')}")
    print(f"사고건수     : {display('occrrnc_cnt')}")
    print(f"사상자수     : {display('caslt_cnt')}")
    print(f"사망자수     : {display('dth_dnv_cnt')}")
    print(f"중상자수     : {display('se_dnv_cnt')}")
    print(f"경상자수     : {display('sl_dnv_cnt')}")
    print(f"부상신고     : {display('wnd_dnv_cnt')}")
    print(f"경도         : {display('lo_crd')}")
    print(f"위도         : {display('la_crd')}")
    print(
        "Polygon      : "
        + ("있음" if not pd.isna(first.get("geom_json")) else "없음")
    )


# ============================================================
# 9. MAIN
# ============================================================


def main():
    started_at = time.monotonic()
    run_id = datetime.now().strftime("%Y%m%d%H%M%S")

    if not ENV_PATH.exists():
        raise FileNotFoundError(
            ".env 파일을 찾을 수 없습니다.\n"
            f"확인 경로: {ENV_PATH}"
        )

    if not API_KEY:
        raise ValueError(
            "BICYCLE_ACCIDENT_API_KEY가 .env에 없습니다."
        )

    RAW_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 70)
    print("KoROAD 자전거 교통사고 전국·전연도 파이프라인")
    print("=" * 70)
    print(f"ENV          : {ENV_PATH}")
    print(f"코드표       : {CODELIST_PATH}")
    print(f"체크포인트 폴더: {RAW_DIR}")
    print("API KEY      : 확인됨")

    scopes, afos_id_by_year = load_request_scopes(CODELIST_PATH)
    print(
        f"수집 범위    : {min(afos_id_by_year)}~"
        f"{max(afos_id_by_year)} / {len(scopes):,}개 요청"
    )
    print()

    raw_df = collect_all_scopes(scopes, run_id)
    display_first_record(raw_df)

    replace_raw_snapshot(raw_df)
    remove_checkpoint_after_success()

    elapsed_seconds = time.monotonic() - started_at

    print()
    print("=" * 70)
    print("전국·전연도 RAW 수집 성공")
    print("=" * 70)
    print(f"요청 범위       : {len(scopes):,}개")
    print(f"수집 건수       : {len(raw_df):,}건")
    print("저장 RAW DB     : raw.koroad_bicycle_accident_hotspots")
    print(f"소요시간        : {elapsed_seconds:,.1f}초")
    print("=" * 70)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n사용자가 실행을 중단했습니다. 기존 DB는 교체하지 않습니다.")
        raise SystemExit(130)
    except Exception as exc:
        print()
        print("=" * 70)
        print("파이프라인 실패")
        print("기존 정상 RAW DB 테이블은 교체하지 않습니다.")
        print(f"원인: {type(exc).__name__}: {exc}")
        print("=" * 70)
        raise
